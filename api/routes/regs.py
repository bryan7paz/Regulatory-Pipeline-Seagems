"""Routes for regulations: CRUD, search, filters, batch validate, export."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session

from ..database import get_db
from .. import models
from ..schemas import AnalysisOut, AnalysisDetail, ValidationAction

router = APIRouter(prefix="/regs", tags=["regs"])


# ── List with search + filters + pagination ────────────────────────

@router.get("/")
def list_analyses(
    q: Optional[str] = None,
    assunto: Optional[str] = None,
    aplicacao: Optional[str] = None,
    status: Optional[str] = None,
    validacao: Optional[str] = None,
    fonte: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    sort_by: str = Query(default="created_at", regex="^(created_at|data_publicacao|entrada_em_vigor|norma|assunto|aplicacao|status|validacao)$"),
    sort_order: str = Query(default="desc", regex="^(asc|desc)$"),
    fields: Optional[str] = None,
    page: int = Query(default=1, ge=1),
    size: int = Query(default=20, ge=1, le=200),
    db: Session = Depends(get_db),
):
    q_query = db.query(models.RegulatoryAnalysis)

    # Full-text search
    if q:
        pattern = f"%{q}%"
        q_query = q_query.filter(
            or_(
                models.RegulatoryAnalysis.norma.ilike(pattern),
                models.RegulatoryAnalysis.requisito.ilike(pattern),
                models.RegulatoryAnalysis.item.ilike(pattern),
                models.RegulatoryAnalysis.itens_modificados.ilike(pattern),
                models.RegulatoryAnalysis.acao_sugerida.ilike(pattern),
                models.RegulatoryAnalysis.source_name.ilike(pattern),
            )
        )

    # Enum filters
    if assunto:
        q_query = q_query.filter(models.RegulatoryAnalysis.assunto == assunto)
    if aplicacao:
        q_query = q_query.filter(models.RegulatoryAnalysis.aplicacao == aplicacao)
    if status:
        q_query = q_query.filter(models.RegulatoryAnalysis.status == status)
    if validacao:
        q_query = q_query.filter(models.RegulatoryAnalysis.status_validacao == validacao)
    if fonte:
        q_query = q_query.filter(models.RegulatoryAnalysis.source_id == fonte)

    # Date range filters
    if date_from:
        try:
            from datetime import date
            df = date.fromisoformat(date_from)
            q_query = q_query.filter(models.RegulatoryAnalysis.data_publicacao >= df)
        except ValueError:
            pass
    if date_to:
        try:
            from datetime import date
            dt = date.fromisoformat(date_to)
            q_query = q_query.filter(models.RegulatoryAnalysis.data_publicacao <= dt)
        except ValueError:
            pass

    # Sorting
    sort_column = getattr(models.RegulatoryAnalysis, sort_by, models.RegulatoryAnalysis.created_at)
    if sort_order == "asc":
        q_query = q_query.order_by(sort_column.asc())
    else:
        q_query = q_query.order_by(sort_column.desc())

    # Pagination
    total = q_query.count()
    offset = (page - 1) * size
    items = q_query.offset(offset).limit(size).all()

    # Field selection
    result_items = [AnalysisOut.model_validate(r) for r in items]
    if fields:
        field_list = [f.strip() for f in fields.split(",")]
        result_items = [
            {k: v for k, v in item.model_dump().items() if k in field_list}
            for item in result_items
        ]

    return {
        "items": result_items,
        "total": total,
        "page": page,
        "size": size,
        "pages": max(1, (total + size - 1) // size),
        "sort_by": sort_by,
        "sort_order": sort_order,
    }


# ── CSV export (MUST be before /{analysis_id}) ────────────────────

@router.get("/export/csv")
def export_csv(
    validacao: Optional[str] = None,
    db: Session = Depends(get_db),
):
    q = db.query(models.RegulatoryAnalysis)
    if validacao:
        q = q.filter(models.RegulatoryAnalysis.status_validacao == validacao)
    rows = q.order_by(models.RegulatoryAnalysis.created_at.desc()).all()

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "id", "fonte", "data_publicacao", "entrada_em_vigor", "requisito",
        "norma", "assunto", "aplicacao", "status", "item", "itens_modificados",
        "acao_sugerida", "validacao", "validado_por", "url_origem",
    ])
    for r in rows:
        writer.writerow([
            r.id, r.source_id,
            r.data_publicacao.isoformat() if r.data_publicacao else "",
            r.entrada_em_vigor.isoformat() if r.entrada_em_vigor else "",
            r.requisito or "", r.norma or "", r.assunto or "",
            r.aplicacao or "", r.status or "", r.item or "",
            r.itens_modificados or "", r.acao_sugerida or "",
            r.status_validacao or "", r.validated_by or "",
            r.url_origem or "",
        ])

    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=regulatory_analysis.csv"},
    )


# ── Batch validate ─────────────────────────────────────────────────

class BatchValidationRequest(BaseModel):
    ids: list[str]
    action: str  # "aprovado" | "reprovado"
    validated_by: str = "especialista"


@router.post("/batch-validate")
def batch_validate(body: BatchValidationRequest, db: Session = Depends(get_db)):
    from core.notify import notify_validation

    if body.action not in ("aprovado", "reprovado"):
        raise HTTPException(status_code=400, detail="action deve ser 'aprovado' ou 'reprovado'")

    rows = (
        db.query(models.RegulatoryAnalysis)
        .filter(models.RegulatoryAnalysis.id.in_(body.ids))
        .all()
    )
    now = datetime.now(timezone.utc)
    updated = 0
    for row in rows:
        row.status_validacao = body.action
        row.validated_at = now
        row.validated_by = body.validated_by[:100]
        updated += 1

    db.commit()

    # Notify batch validation
    if updated > 0:
        notify_validation(
            doc_id=f"batch ({updated} documentos)",
            action=body.action,
            validated_by=body.validated_by,
        )
    return {"updated": updated}


# ── Single detail ──────────────────────────────────────────────────

@router.get("/{analysis_id}")
def get_analysis(analysis_id: str, db: Session = Depends(get_db)):
    row = db.get(models.RegulatoryAnalysis, analysis_id)
    if not row:
        raise HTTPException(status_code=404, detail="Registro não encontrado")
    return AnalysisDetail.model_validate(row)


# ── Validate single ────────────────────────────────────────────────

@router.post("/{analysis_id}/validate", response_model=AnalysisOut)
def validate_single(
    analysis_id: str,
    body: ValidationAction,
    db: Session = Depends(get_db),
):
    from core.notify import notify_validation

    row = db.get(models.RegulatoryAnalysis, analysis_id)
    if not row:
        raise HTTPException(status_code=404, detail="Registro não encontrado")

    row.status_validacao = body.action
    row.validated_at = datetime.now(timezone.utc)
    row.validated_by = (body.validated_by or "especialista")[:100]
    db.commit()
    db.refresh(row)

    notify_validation(analysis_id, body.action, row.validated_by)
    return row


# ── PDF Export ───────────────────────────────────────────────────

@router.get("/export/pdf")
def export_pdf(
    validacao: Optional[str] = None,
    fonte: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """Export analysis results as a formatted PDF."""
    from pathlib import Path
    from jinja2 import Environment, FileSystemLoader
    from weasyprint import HTML

    q = db.query(models.RegulatoryAnalysis)
    if validacao:
        q = q.filter(models.RegulatoryAnalysis.status_validacao == validacao)
    if fonte:
        q = q.filter(models.RegulatoryAnalysis.source_id == fonte)
    rows = q.order_by(models.RegulatoryAnalysis.created_at.desc()).all()

    items = [AnalysisOut.model_validate(r).model_dump() for r in rows]

    # Count statuses
    approved = sum(1 for r in items if r.get("status_validacao") == "aprovado")
    rejected = sum(1 for r in items if r.get("status_validacao") == "reprovado")
    pending = sum(1 for r in items if r.get("status_validacao") in (None, "pendente"))

    # Build filter description
    filters_parts = []
    if validacao:
        filters_parts.append(f"Validação: {validacao}")
    if fonte:
        filters_parts.append(f"Fonte: {fonte}")
    filters = " | ".join(filters_parts) if filters_parts else "Nenhum filtro"

    # Render template
    template_dir = Path(__file__).parent.parent / "templates"
    env = Environment(loader=FileSystemLoader(str(template_dir)))
    template = env.get_template("pdf_report.html")

    html_content = template.render(
        items=items,
        total=len(items),
        approved=approved,
        rejected=rejected,
        pending=pending,
        filters=filters,
        generated_at=datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC"),
    )

    # Generate PDF
    pdf_bytes = HTML(string=html_content).write_pdf()

    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=regulatory_analysis.pdf"},
    )


# ── Excel Export ─────────────────────────────────────────────────

@router.get("/export/excel")
def export_excel(
    validacao: Optional[str] = None,
    fonte: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """Export analysis results as an Excel spreadsheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    q = db.query(models.RegulatoryAnalysis)
    if validacao:
        q = q.filter(models.RegulatoryAnalysis.status_validacao == validacao)
    if fonte:
        q = q.filter(models.RegulatoryAnalysis.source_id == fonte)
    rows = q.order_by(models.RegulatoryAnalysis.created_at.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Análise Regulatória"

    # Headers
    headers = [
        "ID", "Fonte", "Data Publicação", "Entrada em Vigor",
        "Requisito", "Norma", "Assunto", "Aplicação", "Status",
        "Item", "Itens Modificados", "Ação Sugerida", "Validação",
        "Validado por", "URL Origem",
    ]

    # Style headers
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data rows
    for row_idx, row in enumerate(rows, 2):
        data = [
            row.id, row.source_id,
            row.data_publicacao.isoformat() if row.data_publicacao else "",
            row.entrada_em_vigor.isoformat() if row.entrada_em_vigor else "",
            row.requisito or "", row.norma or "", row.assunto or "",
            row.aplicacao or "", row.status or "", row.item or "",
            row.itens_modificados or "", row.acao_sugerida or "",
            row.status_validacao or "", row.validated_by or "",
            row.url_origem or "",
        ]
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Auto-width columns
    for col in range(1, len(headers) + 1):
        max_length = max(
            len(str(ws.cell(row=r, column=col).value or ""))
            for r in range(1, min(ws.max_row + 1, 50))
        )
        ws.column_dimensions[get_column_letter(col)].width = min(max_length + 2, 40)

    # Freeze header
    ws.freeze_panes = "A2"

    # Save to buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=regulatory_analysis.xlsx"},
    )
